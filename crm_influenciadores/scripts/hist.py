import openpyxl,json,os,re,unicodedata
from collections import defaultdict
D=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
def slug(s):
    s=unicodedata.normalize('NFD',str(s).lower()); s=''.join(c for c in s if unicodedata.category(c)!='Mn')
    return re.sub(r'[^a-z0-9._-]+','-',s).strip('-')[:120]
wb=openpyxl.load_workbook(f'{D}/src/base_unica.xlsx',read_only=True,data_only=True)
B=list(wb['BASE'].iter_rows(values_only=True)); H=B[0]; ix={h:i for i,h in enumerate(H)}
def g(r,k): v=r[ix[k]]; return None if v in (None,'') else v
RANK={'Publicado':4,'Pago':5,'Kit enviado':3,'Aprovado':2,'Recusado':1,'Na lista':0}
def st(s):
    s=str(s or '').strip().lower()
    if re.search(r'post|^sim$',s) and not re.search(r'n[aã]o',s): return 'Publicado'
    if re.search(r'contrat|assinad',s): return 'Pago'
    if re.search(r'^\.$|envia|entreg',s): return 'Kit enviado'
    if re.search(r'aprov|confirm|ok|ativo',s): return 'Aprovado'
    if re.search(r'n[aã]o|recus|sem retorno|negou',s): return 'Recusado'
    return 'Na lista'
hist=defaultdict(dict)
for r in wb['HISTÓRICO'].iter_rows(min_row=2,values_only=True):
    p,marca,ano=r[0],r[2],r[3]
    if not p or not marca or str(marca).startswith('YOUPIX'): continue
    h=slug(str(p).lstrip('@'))
    for m in str(marca).split(' | '):
        k=(m.strip(),str(ano or ''))
        s=st(r[6]); 
        if RANK[s]>=RANK.get(hist[h].get(k),-1): hist[h][k]=s
items=[]
for r in B[1:]:
    p=g(r,'PERFIL (@)')
    nome=g(r,'NOME'); email=g(r,'E-MAIL')
    if not p and not email: continue
    h=slug(str(p).lstrip('@')) if p else slug('email-'+str(email).split('|')[0].strip())
    seg=g(r,'SEGUIDORES'); eng=g(r,'ENGAJ.')
    try: seg=int(float(seg)) if seg else None
    except: seg=None
    try: eng=round(float(eng)*100,2) if eng not in (None,) and float(eng)<1 else (round(float(eng),2) if eng else None)
    except: eng=None
    hs=sorted(hist.get(h,{}).items(),key=lambda kv:(kv[0][1],kv[0][0]),reverse=True)[:25]
    it=dict(h=h,perfil=('@'+str(p).lstrip('@')) if p else '',nome=nome or '',email=str(email or '')[:120],tel=str(g(r,'TELEFONE') or '')[:60],ag=str(g(r,'AGÊNCIA / ASSESSORIA') or '')[:80],
            cid=' / '.join(x for x in [str(g(r,'CIDADE') or ''),str(g(r,'UF') or '')] if x),nicho=str(g(r,'NICHO') or '')[:80],seg=seg,eng=eng,faixa=g(r,'FAIXA') or '',
            marcas=str(g(r,'MARCAS (clientes MOVE)') or '')[:200],a1=str(g(r,'1º ANO') or ''),a2=str(g(r,'ÚLT. ANO') or ''),rel=str(g(r,'RELAÇÃO COM A MOVE') or '')[:120],
            cache=g(r,'ÚLT. CACHÊ (R$)'),ln=str(g(r,'LISTA NEGRA') or '').upper().startswith('SIM'),
            hs=[[m,a,s] for (m,a),s in hs])
    for k in list(it):
        if it[k] in ('',None,[],False) and k not in ('h',): del it[k]
    items.append(it)
# dedupe h
seen={};
for it in items:
    if it['h'] in seen: 
        o=seen[it['h']]; o.setdefault('hs',[]); continue
    seen[it['h']]=it
items=list(seen.values())
os.makedirs(f'{D}/s3',exist_ok=True)
chunks=[];cur=[];size=0
for it in items:
    s=len(json.dumps(it,ensure_ascii=False))
    if size+s>180000: chunks.append(cur);cur=[];size=0
    cur.append(it);size+=s
if cur: chunks.append(cur)
W=[]
for i,c in enumerate(chunks):
    p=f'{D}/s3/b{i:03d}.json'; json.dump({'n':len(c),'items':c,'fonte':'Move · Base Única de Influenciadores 2016-2026 (08/10/2026)'},open(p,'w'),ensure_ascii=False)
    W.append({'op':'set','collection':'base_historica','doc_id':f'b{i:03d}','file_path':p})
json.dump(W,open(f'{D}/s3/writes.json','w'))
print(len(items),'perfis',len(chunks),'blocos',max(os.path.getsize(f'{D}/s3/b{i:03d}.json') for i in range(len(chunks))),'bytes max')
print(sum(1 for i in items if i.get('ln')),'lista negra')
print(json.dumps(items[0],ensure_ascii=False)[:600])
