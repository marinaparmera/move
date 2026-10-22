import openpyxl,re,json,sys
from collections import Counter
SRC=sys.argv[1]
FILES={'ST. DALFOUR':'stdalfour','KALASSI':'kalassi_ctrl','SPLENDA':'splenda','MAPED':'maped','JANOD':'janod'}
MES=['Jan','Fev','Mar','Abr','Mai','Jun','Jul','Ago','Set','Out','Nov','Dez']
def tabmes(s):
    m=re.match(r'\s*(Prescritores\s+)?(\d{2})(\d{2})',s)
    if not m: return None
    mm,yy=int(m.group(2)),int(m.group(3))
    if yy<26 and not (yy==25 and False): return None
    if not 1<=mm<=12: return None
    return f'{MES[mm-1]}/{yy}', bool(m.group(1))
out=[]
for marca,f in FILES.items():
    wb=openpyxl.load_workbook(f'{SRC}/{f}.xlsx',read_only=True,data_only=True)
    for s in wb.sheetnames:
        tm=tabmes(s)
        if not tm or int(tm[0][-2:])<26: continue
        mes,presc=tm
        rows=list(wb[s].iter_rows(values_only=True))
        hi=None
        for i,r in enumerate(rows[:6]):
            if any(str(c).strip().upper()=='PERFIL' for c in r if c): hi=i;break
        if hi is None: continue
        H={str(c).strip().upper():j for j,c in enumerate(rows[hi]) if c}
        def g(r,k):
            j=H.get(k); return r[j] if j is not None and j<len(r) else None
        for r in rows[hi+1:]:
            p=g(r,'PERFIL')
            if not p or not str(p).strip(): continue
            p=str(p).strip()
            if len(p)>60 or ' ' in p.strip('@ ') and not p.startswith('@'): 
                # probably a note row
                if len(p.split())>3: continue
            out.append(dict(marca=marca,mes=mes,tab=s,prescritor=presc,perfil=p,seg=g(r,'SEGUIDORES'),eng=g(r,'ENGAJAMENTO'),nicho=g(r,'NICHO'),just=g(r,'JUSTIFICATIVA'),status=g(r,'STATUS'),entrega=g(r,'ENTREGA')))
json.dump(out,open(f'{SRC}/../ctrl_rows.json','w'),default=str,ensure_ascii=False)
print(len(out))
print(Counter((o['marca'],o['mes'],o['prescritor']) for o in out))
print(Counter(str(o['status']).strip().upper()[:40] for o in out).most_common(40))
print(Counter(str(o['entrega']).strip().upper()[:40] for o in out).most_common(20))
